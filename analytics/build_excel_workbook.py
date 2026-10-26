"""Build analytics/claims_analysis.xlsx from the exported marts.

Every summary number is an Excel formula (SUMIFS / INDEX style) over the data
sheets, so the workbook recalculates if the data sheets are refreshed.

Usage (inside the analytics container, after `python -m pipeline.export_bi`):
    python analytics/build_excel_workbook.py
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
NAVY = "1F3864"
LIGHT = "D9E2F3"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

USD = '$#,##0;($#,##0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
YEARS = [2008, 2009, 2010]
LOS_ORDER = ["0-1 days", "2-3 days", "4-7 days", "8-14 days", "15+ days"]

DATA_SHEETS = {
    # sheet name: (csv file, {column: number format})
    "Data_Monthly": ("monthly_claims.csv", {"claims": INT, "patients": INT, "total_payment": USD, "adjustments": INT}),
    "Data_Diagnosis": ("claims_by_diagnosis_chapter.csv", {"claims": INT, "total_payment": USD}),
    "Data_States": ("claims_by_state.csv", {"claims": INT, "patients": INT, "total_payment": USD}),
    "Data_LOS": ("inpatient_length_of_stay.csv", {"claims": INT, "total_payment": USD}),
    "Data_Conditions": ("chronic_conditions.csv", {
        "total_patients": INT, "patients_with_condition": INT,
        "reimbursement_with_condition": USD, "reimbursement_without_condition": USD,
    }),
}


def font(**kw) -> Font:
    return Font(name=FONT, **kw)


def col_range(sheet: str, df: pd.DataFrame, column: str) -> str:
    """Absolute range of one data column, e.g. Data_Monthly!$F$2:$F$76."""
    letter = get_column_letter(df.columns.get_loc(column) + 1)
    return f"{sheet}!${letter}$2:${letter}${len(df) + 1}"


def title(ws, text: str, subtitle: str) -> None:
    ws["B2"] = text
    ws["B2"].font = font(size=16, bold=True, color=NAVY)
    ws["B3"] = subtitle
    ws["B3"].font = font(size=10, italic=True, color="595959")
    ws.column_dimensions["A"].width = 2


def header_row(ws, row: int, col: int, labels: list[str]) -> None:
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=col + i, value=label)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30


def body_cell(ws, row: int, col: int, value, fmt: str | None = None, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font(bold=bold)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    return cell


def add_data_sheet(wb: Workbook, name: str, df: pd.DataFrame, formats: dict) -> None:
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for record in df.itertuples(index=False):
        ws.append(list(record))
    for c, column in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c).font = font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = max(14, len(column) + 4)
        for r in range(2, len(df) + 2):
            cell = ws.cell(row=r, column=c)
            cell.font = font()
            if column in formats:
                cell.number_format = formats[column]
    table = Table(displayName=name, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

def summary_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Summary", 0)
    title(ws, "Medicare Claims Analysis, 2008-2010",
          "CMS DE-SynPUF synthetic Medicare data (Sample 1). Source: gold layer of the healthcare-claims-etl pipeline.")
    m = data["Data_Monthly"]
    m_year, m_type = col_range("Data_Monthly", m, "claim_year"), col_range("Data_Monthly", m, "claim_type")
    m_claims, m_pay = col_range("Data_Monthly", m, "claims"), col_range("Data_Monthly", m, "total_payment")

    # KPI tiles (2008-2010; the 536 claims that start in 2007 are left out)
    kpis = [
        ("B", "Total Medicare payment", f'=SUMIFS({m_pay},{m_year},">=2008")', USD),
        ("D", "Claims", f'=SUMIFS({m_claims},{m_year},">=2008")', INT),
        ("F", "Average payment per claim", "=IF(D6=0,0,B6/D6)", '$#,##0'),
        ("H", "Inpatient share of payment", f'=IF(B6=0,0,SUMIFS({m_pay},{m_year},">=2008",{m_type},"IP")/B6)', PCT),
    ]
    for col, label, formula, fmt in kpis:
        ws[f"{col}5"] = label
        ws[f"{col}5"].font = font(size=9, color="595959")
        ws[f"{col}6"] = formula
        ws[f"{col}6"].font = font(size=14, bold=True, color=NAVY)
        ws[f"{col}6"].number_format = fmt
    ws["B7"] = "Years 2008-2010. The 536 claims that start in 2007 are excluded."
    ws["B7"].font = font(size=8, italic=True, color="7F7F7F")

    # Table A: payment by year and claim type
    ws["B9"] = "Payment and claims by year"
    ws["B9"].font = font(size=12, bold=True, color=NAVY)
    header_row(ws, 10, 2, ["Year", "Inpatient claims", "Inpatient payment", "Outpatient claims",
                           "Outpatient payment", "Total payment", "Change vs prior year"])
    for i, year in enumerate(YEARS):
        r = 11 + i
        body_cell(ws, r, 2, year, "0", bold=True)
        body_cell(ws, r, 3, f'=SUMIFS({m_claims},{m_year},$B{r},{m_type},"IP")', INT)
        body_cell(ws, r, 4, f'=SUMIFS({m_pay},{m_year},$B{r},{m_type},"IP")', USD)
        body_cell(ws, r, 5, f'=SUMIFS({m_claims},{m_year},$B{r},{m_type},"OP")', INT)
        body_cell(ws, r, 6, f'=SUMIFS({m_pay},{m_year},$B{r},{m_type},"OP")', USD)
        body_cell(ws, r, 7, f"=D{r}+F{r}", USD)
        body_cell(ws, r, 8, "-" if i == 0 else f'=IF(G{r - 1}=0,"-",G{r}/G{r - 1}-1)', PCT)
    total = 11 + len(YEARS)
    body_cell(ws, total, 2, "Total", bold=True)
    for c in range(3, 8):
        letter = get_column_letter(c)
        body_cell(ws, total, c, f"=SUM({letter}11:{letter}{total - 1})", INT if c in (3, 5) else USD, bold=True)
    body_cell(ws, total, 8, "")
    ws.cell(row=total + 1, column=2, value="The CMS sample has fewer 2010 claims than earlier years, so 2010 totals are lower.").font = \
        font(size=8, italic=True, color="7F7F7F")

    # Table C: top diagnosis chapters
    d = data["Data_Diagnosis"]
    d_year, d_chapter = col_range("Data_Diagnosis", d, "claim_year"), col_range("Data_Diagnosis", d, "icd9_chapter")
    d_claims, d_pay = col_range("Data_Diagnosis", d, "claims"), col_range("Data_Diagnosis", d, "total_payment")
    ranked = (d[d["claim_year"] >= 2008].groupby("icd9_chapter")["total_payment"].sum()
              .sort_values(ascending=False).index.tolist())
    top = ranked[:10]

    start = total + 4
    ws.cell(row=start - 1, column=2, value="Top 10 diagnosis chapters by payment (primary diagnosis, 2008-2010)").font = \
        font(size=12, bold=True, color=NAVY)
    header_row(ws, start, 2, ["ICD-9 chapter", "Claims", "Payment", "Share of payment", "Average payment per claim"])
    for i, chapter in enumerate(top):
        r = start + 1 + i
        body_cell(ws, r, 2, chapter)
        body_cell(ws, r, 3, f'=SUMIFS({d_claims},{d_year},">=2008",{d_chapter},$B{r})', INT)
        body_cell(ws, r, 4, f'=SUMIFS({d_pay},{d_year},">=2008",{d_chapter},$B{r})', USD)
        body_cell(ws, r, 5, f"=IF($B$6=0,0,D{r}/$B$6)", PCT)
        body_cell(ws, r, 6, f"=IF(C{r}=0,0,D{r}/C{r})", '$#,##0')
    other = start + 1 + len(top)
    body_cell(ws, other, 2, "All other chapters")
    body_cell(ws, other, 3, f"=$D$6-SUM(C{start + 1}:C{other - 1})", INT)
    body_cell(ws, other, 4, f"=$B$6-SUM(D{start + 1}:D{other - 1})", USD)
    body_cell(ws, other, 5, f"=IF($B$6=0,0,D{other}/$B$6)", PCT)
    body_cell(ws, other, 6, f"=IF(C{other}=0,0,D{other}/C{other})", '$#,##0')
    ws.cell(row=other + 1, column=2,
            value="Chapters are ranked by 2008-2010 payment when the workbook is built; values are live formulas.").font = \
        font(size=8, italic=True, color="7F7F7F")

    widths = {"B": 44, "C": 15, "D": 18, "E": 15, "F": 18, "G": 18, "H": 16}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Payment by year and claim type"
    chart.y_axis.title = "USD"
    chart.y_axis.numFmt = '$#,##0,,"M"'
    chart.add_data(Reference(ws, min_col=4, min_row=10, max_row=10 + len(YEARS)), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=6, min_row=10, max_row=10 + len(YEARS)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=11, max_row=10 + len(YEARS)))
    chart.height, chart.width = 7.5, 13
    ws.add_chart(chart, "J9")

    bars = BarChart()
    bars.type = "bar"
    bars.title = "Top diagnosis chapters by payment"
    bars.legend = None
    bars.x_axis.scaling.orientation = "maxMin"
    bars.y_axis.numFmt = '$#,##0,,"M"'
    bars.add_data(Reference(ws, min_col=4, min_row=start, max_row=start + len(top)), titles_from_data=True)
    bars.set_categories(Reference(ws, min_col=2, min_row=start + 1, max_row=start + len(top)))
    bars.height, bars.width = 9, 13
    ws.add_chart(bars, f"J{start - 1}")
    ws.freeze_panes = "A8"


# ---------------------------------------------------------------------------
# Monthly trend
# ---------------------------------------------------------------------------

def trend_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Monthly Trend", 1)
    title(ws, "Monthly payment trend", "Medicare payment per month by claim type, 2008-2010.")
    m = data["Data_Monthly"]
    m_year, m_month = col_range("Data_Monthly", m, "claim_year"), col_range("Data_Monthly", m, "month")
    m_type, m_pay = col_range("Data_Monthly", m, "claim_type"), col_range("Data_Monthly", m, "total_payment")
    m_claims = col_range("Data_Monthly", m, "claims")

    header_row(ws, 5, 2, ["Year", "Month", "Month label", "Inpatient payment", "Outpatient payment",
                          "Total payment", "Claims"])
    r = 6
    for year in YEARS:
        for month in range(1, 13):
            body_cell(ws, r, 2, year, "0")
            body_cell(ws, r, 3, month, "0")
            body_cell(ws, r, 4, f'=B{r}&"-"&TEXT(C{r},"00")')
            body_cell(ws, r, 5, f'=SUMIFS({m_pay},{m_year},$B{r},{m_month},$C{r},{m_type},"IP")', USD)
            body_cell(ws, r, 6, f'=SUMIFS({m_pay},{m_year},$B{r},{m_month},$C{r},{m_type},"OP")', USD)
            body_cell(ws, r, 7, f"=E{r}+F{r}", USD)
            body_cell(ws, r, 8, f"=SUMIFS({m_claims},{m_year},$B{r},{m_month},$C{r})", INT)
            r += 1
    last = r - 1
    for letter, width in {"B": 8, "C": 8, "D": 12, "E": 18, "F": 18, "G": 18, "H": 12}.items():
        ws.column_dimensions[letter].width = width

    chart = LineChart()
    chart.title = "Monthly Medicare payment"
    chart.y_axis.title = "USD"
    chart.y_axis.numFmt = '$#,##0,,"M"'
    chart.add_data(Reference(ws, min_col=5, max_col=6, min_row=5, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=4, min_row=6, max_row=last))
    chart.height, chart.width = 9, 22
    ws.add_chart(chart, "J5")
    ws.freeze_panes = "A6"


# ---------------------------------------------------------------------------
# Chronic conditions (interactive year)
# ---------------------------------------------------------------------------

def conditions_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Chronic Conditions", 2)
    title(ws, "Chronic conditions: prevalence and cost",
          "Share of patients with each condition and their average annual Medicare reimbursement.")
    c = data["Data_Conditions"]
    c_year, c_cond = col_range("Data_Conditions", c, "coverage_year"), col_range("Data_Conditions", c, "condition")
    c_total = col_range("Data_Conditions", c, "total_patients")
    c_with = col_range("Data_Conditions", c, "patients_with_condition")
    c_rw = col_range("Data_Conditions", c, "reimbursement_with_condition")
    c_rwo = col_range("Data_Conditions", c, "reimbursement_without_condition")

    ws["B5"] = "Coverage year (edit):"
    ws["B5"].font = font(bold=True)
    ws["C5"] = 2010
    ws["C5"].fill = INPUT_FILL
    ws["C5"].font = font(bold=True, color="0000FF")
    ws["C5"].number_format = "0"
    ws["C5"].comment = Comment("Enter 2008, 2009 or 2010. Every number on this sheet updates.", "analytics")
    ws["D5"] = "Yellow cell is the only input on this sheet."
    ws["D5"].font = font(size=8, italic=True, color="7F7F7F")

    header_row(ws, 7, 2, ["Condition", "Patients with condition", "Total patients", "Prevalence",
                          "Avg reimbursement, with condition", "Avg reimbursement, without", "Cost multiple"])
    conditions = sorted(c["condition"].unique())
    for i, condition in enumerate(conditions):
        r = 8 + i
        body_cell(ws, r, 2, condition)
        body_cell(ws, r, 3, f"=SUMIFS({c_with},{c_year},$C$5,{c_cond},$B{r})", INT)
        body_cell(ws, r, 4, f"=SUMIFS({c_total},{c_year},$C$5,{c_cond},$B{r})", INT)
        body_cell(ws, r, 5, f"=IF(D{r}=0,0,C{r}/D{r})", PCT)
        body_cell(ws, r, 6, f"=IF(C{r}=0,0,SUMIFS({c_rw},{c_year},$C$5,{c_cond},$B{r})/C{r})", '$#,##0')
        body_cell(ws, r, 7, f"=IF(D{r}-C{r}=0,0,SUMIFS({c_rwo},{c_year},$C$5,{c_cond},$B{r})/(D{r}-C{r}))", '$#,##0')
        body_cell(ws, r, 8, f'=IF(G{r}=0,0,F{r}/G{r})', '0.0"x"')
    last = 7 + len(conditions)
    ws.cell(row=last + 2, column=2, value=(
        "Reimbursement = annual Medicare inpatient + outpatient + carrier reimbursement from the beneficiary file. "
        "Cost multiple = average for patients with the condition / average for patients without it."
    )).font = font(size=8, italic=True, color="7F7F7F")
    for letter, width in {"B": 30, "C": 14, "D": 14, "E": 12, "F": 20, "G": 20, "H": 12}.items():
        ws.column_dimensions[letter].width = width

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Prevalence by condition (selected year)"
    chart.legend = None
    chart.y_axis.numFmt = "0%"
    chart.add_data(Reference(ws, min_col=5, min_row=7, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=8, max_row=last))
    chart.height, chart.width = 9, 14
    ws.add_chart(chart, "J7")


# ---------------------------------------------------------------------------
# Inpatient length of stay
# ---------------------------------------------------------------------------

def los_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Inpatient Stays", 3)
    title(ws, "Inpatient length of stay", "Inpatient claims and payment by length-of-stay band.")
    l = data["Data_LOS"]
    l_year, l_band = col_range("Data_LOS", l, "claim_year"), col_range("Data_LOS", l, "length_of_stay_band")
    l_claims, l_pay = col_range("Data_LOS", l, "claims"), col_range("Data_LOS", l, "total_payment")

    header_row(ws, 5, 2, ["Length of stay", *[f"Claims {y}" for y in YEARS], "Claims 2008-2010",
                          "Payment 2008-2010", "Average payment per stay", "Share of stays"])
    for i, band in enumerate(LOS_ORDER):
        r = 6 + i
        body_cell(ws, r, 2, band)
        for j, year in enumerate(YEARS):
            body_cell(ws, r, 3 + j, f"=SUMIFS({l_claims},{l_year},{year},{l_band},$B{r})", INT)
        body_cell(ws, r, 6, f"=SUM(C{r}:E{r})", INT)
        body_cell(ws, r, 7, f'=SUMIFS({l_pay},{l_year},">=2008",{l_band},$B{r})', USD)
        body_cell(ws, r, 8, f"=IF(F{r}=0,0,G{r}/F{r})", '$#,##0')
        body_cell(ws, r, 9, f"=IF(SUM($F$6:$F${5 + len(LOS_ORDER)})=0,0,F{r}/SUM($F$6:$F${5 + len(LOS_ORDER)}))", PCT)
    last = 5 + len(LOS_ORDER)
    for letter, width in {"B": 16, "C": 12, "D": 12, "E": 12, "F": 16, "G": 18, "H": 18, "I": 12}.items():
        ws.column_dimensions[letter].width = width

    chart = BarChart()
    chart.type = "col"
    chart.title = "Average inpatient payment by length of stay"
    chart.legend = None
    chart.y_axis.numFmt = "$#,##0"
    chart.add_data(Reference(ws, min_col=8, min_row=5, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=6, max_row=last))
    chart.height, chart.width = 8, 14
    ws.add_chart(chart, "K5")


def notes_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Notes")
    title(ws, "Notes", "Where the data comes from and how to use this workbook.")
    lines = [
        ("Source", "CMS 2008-2010 Data Entrepreneurs' Synthetic Public Use File (DE-SynPUF), Sample 1: synthetic Medicare data, "
                   "https://www.cms.gov/data-research/statistics-trends-and-reports/medicare-claims-synthetic-public-use-files"),
        ("Pipeline", "Data comes from the gold layer of https://github.com/siddhidokania/healthcare-claims-etl "
                     "(python -m pipeline.export_bi writes the Data_ sheets' CSVs)."),
        ("Claims", "One row per claim; claims split into two source segments are consolidated. 278 invalid claim segments "
                   "were quarantined by the pipeline and are not included."),
        ("Payment", "Medicare claim payment amount (CLM_PMT_AMT) in USD. Negative values are adjustments."),
        ("Diagnosis chapter", "ICD-9-CM chapter of the claim's primary diagnosis."),
        ("Chronic conditions", "From the yearly beneficiary summary files; reimbursement is annual inpatient + outpatient + carrier."),
        ("Formulas", "All summary numbers are SUMIFS formulas over the Data_ sheets. Replace a Data_ sheet's rows with a "
                     "newer export and the summaries update."),
        ("Input cells", "Yellow cells with blue text are inputs (Chronic Conditions!C5)."),
        ("Practice", "Insert > PivotTable on Data_Monthly or Data_Diagnosis to explore further (e.g. payment by year and chapter)."),
    ]
    for i, (key, text) in enumerate(lines):
        ws.cell(row=5 + i, column=2, value=key).font = font(bold=True)
        cell = ws.cell(row=5 + i, column=3, value=text)
        cell.font = font()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[5 + i].height = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 110


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--marts", default="data/exports/marts")
    parser.add_argument("--out", default="analytics/claims_analysis.xlsx")
    args = parser.parse_args()

    data = {name: pd.read_csv(Path(args.marts) / csv, dtype={"state_code": str})
            for name, (csv, _) in DATA_SHEETS.items()}
    wb = Workbook()
    wb.remove(wb.active)
    for name, (_, formats) in DATA_SHEETS.items():
        add_data_sheet(wb, name, data[name], formats)
    summary_sheet(wb, data)
    trend_sheet(wb, data)
    conditions_sheet(wb, data)
    los_sheet(wb, data)
    notes_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
